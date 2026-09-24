"""
Lógica central da classificação de UPG (Realizados x Compromisso).

Fluxo:
1. `preparar`: le as duas planilhas brutas do SAP, filtra as linhas cuja coluna
   OBJETO contenha "UPG", aplica automaticamente as classificações já
   conhecidas (vindas do "banco" persistente) e gera um Excel de trabalho
   com as linhas novas destacadas para classificação manual.
2. `consolidar`: le o Excel de trabalho já classificado manualmente e grava
   (upsert) as classificações novas/alteradas no banco persistente, para que
   sejam reaproveitadas automaticamente no mês seguinte.

O "banco" é apenas um arquivo Excel (classificacoes.xlsx) com uma linha por
PO. Qualquer pessoa pode abri-lo e conferir/editar diretamente, sem precisar
de Python.
"""

from __future__ import annotations

import datetime as dt
from dataclasses import dataclass
from pathlib import Path
from typing import Callable

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

LogFn = Callable[[str], None]

STATUS_HERDADO = "OK (herdado)"
STATUS_PENDENTE = "PENDENTE"

DB_COLUMNS = ["PO", "Classificacao", "Observacao", "AtualizadoEm"]

ABA_REALIZADOS = "Realizados"
ABA_COMPROMISSO = "Compromisso"

DEFAULT_BANCO = Path("classificacoes.xlsx")

# Nomes reais das colunas de PO em cada relatório do SAP (confirmados pelo usuário).
DEFAULT_PO_COL_REALIZADOS = "Documento de compras"
DEFAULT_PO_COL_COMPROMISSO = "Nº doc.de referência"

PENDENTE_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)


@dataclass
class ReportSpec:
    """Configuração de leitura de um relatório do SAP (Realizados ou Compromisso)."""

    nome: str
    caminho: Path
    coluna_po: str
    coluna_objeto: str = "OBJETO"
    aba: str | int | None = 0
    linha_cabecalho: int = 0


EXTENSOES_EXCEL = (".xlsx", ".xls", ".xlsm")


def localizar_arquivo_por_padrao(pasta: Path, padrao: str, log: LogFn = print) -> Path:
    """Procura, dentro de `pasta`, um arquivo Excel cujo nome contenha
    `padrao` (ex: "Realizado", "Compromisso"), ignorando maiúsculas/
    minúsculas. O nome exportado pelo SAP muda todo mês por causa da data,
    mas a palavra-chave (Realizado/Compromisso) é sempre a mesma.

    Se houver mais de um arquivo correspondente (ex: sobrou um arquivo do
    mês passado na mesma pasta), escolhe o modificado mais recentemente e
    avisa sobre os demais candidatos, para que o usuário possa conferir.
    """

    if not pasta.is_dir():
        raise FileNotFoundError(f"Pasta não encontrada: {pasta}")

    candidatos = sorted(
        (
            arquivo
            for arquivo in pasta.iterdir()
            if arquivo.is_file()
            and not arquivo.name.startswith("~$")  # arquivo temporário do Excel (aberto)
            and arquivo.suffix.lower() in EXTENSOES_EXCEL
            and padrao.lower() in arquivo.name.lower()
        ),
        key=lambda p: p.stat().st_mtime,
        reverse=True,
    )

    if not candidatos:
        raise FileNotFoundError(
            f"Nenhum arquivo contendo '{padrao}' no nome foi encontrado em: {pasta.resolve()}"
        )

    escolhido = candidatos[0]
    if len(candidatos) > 1:
        outros = ", ".join(c.name for c in candidatos[1:])
        log(
            f"Aviso: mais de um arquivo contendo '{padrao}' encontrado em {pasta.resolve()}. "
            f"Usando o mais recente: {escolhido.name} (ignorados: {outros})"
        )
    else:
        log(f"Arquivo de '{padrao}' encontrado: {escolhido.name}")

    return escolhido


def _find_column(columns: list[str], candidates: list[str]) -> str | None:
    """Procura, sem diferenciar maiúsculas/minúsculas, uma coluna cujo nome
    contenha algum dos textos candidatos. Retorna None se não achar exatamente
    uma correspondência (nenhuma ou mais de uma)."""

    normalized = {c: c.strip().lower() for c in columns}
    matches = [
        original
        for original, low in normalized.items()
        if any(cand.lower() in low for cand in candidates)
    ]
    if len(matches) == 1:
        return matches[0]
    return None


def sugerir_coluna_objeto(df: pd.DataFrame) -> str | None:
    return _find_column(list(df.columns), ["objeto"])


def sugerir_coluna_po(df: pd.DataFrame) -> str | None:
    return _find_column(
        list(df.columns),
        ["pedido de compra", "purchase order", "nº pedido", "n° pedido", "pedido", "po"],
    )


def _resolver_nome_coluna(columns: list[str], nome_desejado: str) -> str | None:
    """Encontra o nome exato de uma coluna (preservando a capitalização real
    usada na planilha), sem diferenciar maiúsculas/minúsculas nem espaços nas
    pontas. O SAP não é consistente entre exportações (ex: 'OBJETO' numa
    planilha e 'Objeto' na outra)."""

    alvo = nome_desejado.strip().lower()
    for coluna in columns:
        if coluna.strip().lower() == alvo:
            return coluna
    return None


def carregar_planilha(spec: ReportSpec) -> tuple[pd.DataFrame, str, str]:
    """Lê a planilha e resolve os nomes reais das colunas OBJETO e PO
    (ignorando maiúsculas/minúsculas). Retorna (dataframe, coluna_objeto,
    coluna_po) já com a capitalização exata usada na planilha."""

    if not spec.caminho.exists():
        raise FileNotFoundError(f"Arquivo não encontrado: {spec.caminho}")
    df = pd.read_excel(spec.caminho, sheet_name=spec.aba, header=spec.linha_cabecalho, dtype=object)

    coluna_objeto = _resolver_nome_coluna(list(df.columns), spec.coluna_objeto)
    if coluna_objeto is None:
        sugestao = sugerir_coluna_objeto(df)
        dica = f" Você quis dizer '{sugestao}'?" if sugestao else ""
        raise ValueError(
            f"[{spec.nome}] Coluna OBJETO '{spec.coluna_objeto}' não encontrada.{dica}\n"
            f"Colunas disponíveis: {list(df.columns)}"
        )

    coluna_po = _resolver_nome_coluna(list(df.columns), spec.coluna_po)
    if coluna_po is None:
        sugestao = sugerir_coluna_po(df)
        dica = f" Você quis dizer '{sugestao}'?" if sugestao else ""
        raise ValueError(
            f"[{spec.nome}] Coluna de PO '{spec.coluna_po}' não encontrada.{dica}\n"
            f"Colunas disponíveis: {list(df.columns)}"
        )

    return df, coluna_objeto, coluna_po


def normalizar_po(valor) -> str:
    """Normaliza o número do PO para comparação estável entre planilhas.

    O pandas às vezes lê números de PO como float (ex: 4500123456.0) quando a
    coluna tem células vazias misturadas; aqui garantimos que o mesmo PO gere
    sempre a mesma chave de texto, independente de ter vindo como float, int
    ou string com espaços.
    """

    if valor is None or (isinstance(valor, float) and pd.isna(valor)):
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def filtrar_upg(df: pd.DataFrame, coluna_objeto: str) -> pd.DataFrame:
    mascara = df[coluna_objeto].astype(str).str.contains("UPG", case=False, na=False)
    return df.loc[mascara].copy()


def carregar_banco(caminho: Path) -> pd.DataFrame:
    if caminho.exists():
        banco = pd.read_excel(caminho, dtype=object)
        for coluna in DB_COLUMNS:
            if coluna not in banco.columns:
                banco[coluna] = pd.NA
        banco = banco[DB_COLUMNS]
    else:
        banco = pd.DataFrame(columns=DB_COLUMNS)
    banco["PO"] = banco["PO"].map(normalizar_po)
    banco = banco[banco["PO"] != ""]
    banco = banco.drop_duplicates(subset="PO", keep="last")
    return banco


def aplicar_classificacoes(df: pd.DataFrame, coluna_po: str, banco: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    df["_PO_NORMALIZADO"] = df[coluna_po].map(normalizar_po)
    banco_indexado = banco.set_index("PO")[["Classificacao", "Observacao"]]

    df["Classificacao"] = df["_PO_NORMALIZADO"].map(banco_indexado["Classificacao"])
    df["Observacao"] = df["_PO_NORMALIZADO"].map(banco_indexado["Observacao"])
    df["Status"] = df["Classificacao"].apply(
        lambda v: STATUS_HERDADO if pd.notna(v) and str(v).strip() != "" else STATUS_PENDENTE
    )
    df = df.drop(columns=["_PO_NORMALIZADO"])
    return df


def _formatar_aba(ws: Worksheet, df: pd.DataFrame, coluna_status: str = "Status") -> None:
    for col_idx, nome_coluna in enumerate(df.columns, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        largura = min(max(len(str(nome_coluna)), 10) + 2, 40)
        ws.column_dimensions[get_column_letter(col_idx)].width = largura

    if coluna_status not in df.columns:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        return

    status_idx = list(df.columns).index(coluna_status) + 1
    for row_idx in range(2, ws.max_row + 1):
        if ws.cell(row=row_idx, column=status_idx).value == STATUS_PENDENTE:
            for col_idx in range(1, ws.max_column + 1):
                ws.cell(row=row_idx, column=col_idx).fill = PENDENTE_FILL

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def salvar_planilha_trabalho(
    caminho_saida: Path,
    dados: dict[str, pd.DataFrame],
) -> None:
    with pd.ExcelWriter(caminho_saida, engine="openpyxl") as writer:
        resumo_linhas = []
        for nome_aba, df in dados.items():
            df.to_excel(writer, sheet_name=nome_aba[:31], index=False)
            total = len(df)
            herdadas = int((df["Status"] == STATUS_HERDADO).sum()) if "Status" in df.columns else 0
            pendentes = total - herdadas
            resumo_linhas.append(
                {
                    "Relatorio": nome_aba,
                    "Total_linhas_UPG": total,
                    "Herdadas_do_banco": herdadas,
                    "Pendentes_classificar": pendentes,
                }
            )
        resumo_df = pd.DataFrame(resumo_linhas)
        resumo_df.to_excel(writer, sheet_name="Resumo", index=False)

        for nome_aba, df in dados.items():
            ws = writer.sheets[nome_aba[:31]]
            _formatar_aba(ws, df)
        _formatar_aba(writer.sheets["Resumo"], resumo_df, coluna_status="__nenhuma__")


def consolidar_no_banco(
    caminho_trabalho: Path,
    caminho_banco: Path,
    coluna_po_por_aba: dict[str, str],
) -> tuple[pd.DataFrame, int]:
    """Lê o Excel de trabalho já classificado manualmente e faz upsert das
    classificações preenchidas no banco persistente.

    Retorna o banco atualizado e a quantidade de POs novos/alterados.
    """

    banco = carregar_banco(caminho_banco)
    banco_indexado = banco.set_index("PO") if not banco.empty else pd.DataFrame(columns=DB_COLUMNS[1:]).set_index(
        pd.Index([], name="PO")
    )

    agora = dt.datetime.now().strftime("%Y-%m-%d %H:%M")
    alterados = 0

    for nome_aba, coluna_po in coluna_po_por_aba.items():
        df = pd.read_excel(caminho_trabalho, sheet_name=nome_aba, dtype=object)
        if "Classificacao" not in df.columns:
            continue
        coluna_po_real = _resolver_nome_coluna(list(df.columns), coluna_po)
        if coluna_po_real is None:
            raise ValueError(
                f"[{nome_aba}] Coluna de PO '{coluna_po}' não encontrada na planilha de trabalho.\n"
                f"Colunas disponíveis: {list(df.columns)}"
            )
        for _, linha in df.iterrows():
            po = normalizar_po(linha.get(coluna_po_real))
            classificacao = linha.get("Classificacao")
            if po == "" or pd.isna(classificacao) or str(classificacao).strip() == "":
                continue
            classificacao = str(classificacao).strip()
            observacao = linha.get("Observacao")
            observacao = "" if pd.isna(observacao) else str(observacao).strip()

            existente = banco_indexado.loc[po] if po in banco_indexado.index else None
            mudou = (
                existente is None
                or str(existente.get("Classificacao", "")).strip() != classificacao
                or str(existente.get("Observacao", "")).strip() != observacao
            )
            if mudou:
                banco_indexado.loc[po] = {
                    "Classificacao": classificacao,
                    "Observacao": observacao,
                    "AtualizadoEm": agora,
                }
                alterados += 1

    banco_final = banco_indexado.reset_index().rename(columns={"index": "PO"})
    banco_final = banco_final[DB_COLUMNS]
    banco_final = banco_final.sort_values("PO").reset_index(drop=True)
    return banco_final, alterados


def salvar_banco(banco: pd.DataFrame, caminho: Path) -> None:
    with pd.ExcelWriter(caminho, engine="openpyxl") as writer:
        banco.to_excel(writer, sheet_name="Classificacoes", index=False)
        _formatar_aba(writer.sheets["Classificacoes"], banco, coluna_status="__nenhuma__")


def resolver_caminho(caminho_informado: str | None, pasta: str, padrao_nome: str, log: LogFn = print) -> Path:
    """Se o caminho não foi informado explicitamente, procura na pasta o
    arquivo cujo nome contenha `padrao_nome` (ex: "Realizado"), já que o SAP
    sempre muda o nome do arquivo por causa da data."""

    if caminho_informado:
        return Path(caminho_informado)
    return localizar_arquivo_por_padrao(Path(pasta), padrao_nome, log=log)


def executar_preparar(
    *,
    saida: str,
    pasta: str = ".",
    realizados: str | None = None,
    compromisso: str | None = None,
    po_col_realizados: str = DEFAULT_PO_COL_REALIZADOS,
    po_col_compromisso: str = DEFAULT_PO_COL_COMPROMISSO,
    objeto_col: str = "OBJETO",
    aba_realizados: str | int = 0,
    aba_compromisso: str | int = 0,
    header_realizados: int = 0,
    header_compromisso: int = 0,
    banco: str = str(DEFAULT_BANCO),
    log: LogFn = print,
) -> Path:
    """Fluxo completo do comando 'preparar': localiza os arquivos do SAP,
    filtra as linhas de UPG, aplica as classificações já conhecidas e salva
    o Excel de trabalho do mês. Retorna o caminho do arquivo gerado."""

    banco_df = carregar_banco(Path(banco))

    caminho_realizados = resolver_caminho(realizados, pasta, "Realizado", log=log)
    caminho_compromisso = resolver_caminho(compromisso, pasta, "Compromisso", log=log)

    dados: dict[str, pd.DataFrame] = {}
    especificacoes = [
        (ABA_REALIZADOS, caminho_realizados, po_col_realizados, aba_realizados, header_realizados),
        (ABA_COMPROMISSO, caminho_compromisso, po_col_compromisso, aba_compromisso, header_compromisso),
    ]

    for nome_aba, caminho, coluna_po, aba, linha_cabecalho in especificacoes:
        spec = ReportSpec(
            nome=nome_aba,
            caminho=Path(caminho),
            coluna_po=coluna_po,
            coluna_objeto=objeto_col,
            aba=aba,
            linha_cabecalho=linha_cabecalho,
        )
        df, coluna_objeto_real, coluna_po_real = carregar_planilha(spec)
        df_upg = filtrar_upg(df, coluna_objeto_real)
        df_classificado = aplicar_classificacoes(df_upg, coluna_po_real, banco_df)
        dados[nome_aba] = df_classificado

        total = len(df_classificado)
        herdadas = int((df_classificado["Status"] == STATUS_HERDADO).sum())
        log(f"[{nome_aba}] {total} linhas com UPG | {herdadas} já classificadas | {total - herdadas} pendentes")

    caminho_saida = Path(saida)
    salvar_planilha_trabalho(caminho_saida, dados)
    log(f"Arquivo de trabalho gerado em: {caminho_saida.resolve()}")
    return caminho_saida


def executar_consolidar(
    *,
    trabalho: str,
    po_col_realizados: str = DEFAULT_PO_COL_REALIZADOS,
    po_col_compromisso: str = DEFAULT_PO_COL_COMPROMISSO,
    banco: str = str(DEFAULT_BANCO),
    log: LogFn = print,
) -> tuple[Path, int, int]:
    """Fluxo completo do comando 'consolidar': lê o Excel de trabalho já
    classificado manualmente e grava as classificações novas/alteradas no
    banco. Retorna (caminho_do_banco, quantidade_alterada, total_no_banco)."""

    coluna_po_por_aba = {
        ABA_REALIZADOS: po_col_realizados,
        ABA_COMPROMISSO: po_col_compromisso,
    }
    banco_atualizado, alterados = consolidar_no_banco(
        caminho_trabalho=Path(trabalho),
        caminho_banco=Path(banco),
        coluna_po_por_aba=coluna_po_por_aba,
    )
    salvar_banco(banco_atualizado, Path(banco))
    log(f"Banco atualizado: {Path(banco).resolve()}")
    log(f"{alterados} classificação(ões) nova(s) ou alterada(s) salvas.")
    log(f"Total de POs no banco: {len(banco_atualizado)}")
    return Path(banco), alterados, len(banco_atualizado)
